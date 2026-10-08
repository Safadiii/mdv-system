from io import BytesIO
from datetime import datetime

from sqlmodel import Session, select

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from database.models import (
    Product,
    Supplier,
    Customer,
    Purchase,
    PurchaseItem,
    Sale,
    SaleItem,
    InventoryMovement,
)


# ============================================================
# HELPERS
# ============================================================

def enum_value(value):

    if value is None:
        return None

    if hasattr(value, "value"):
        return value.value

    return value


def format_datetime(value):

    if value is None:
        return None

    if isinstance(value, datetime):

        return value.strftime(
            "%Y-%m-%d %H:%M:%S"
        )

    return value


def style_sheet(sheet):

    # --------------------------------------------------------
    # HEADER STYLE
    # --------------------------------------------------------

    header_fill = PatternFill(
        "solid",
        fgColor="E2E8F0"
    )

    header_font = Font(
        bold=True,
        color="0F172A"
    )


    for cell in sheet[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center"
        )


    # Keep headers visible while scrolling
    sheet.freeze_panes = "A2"


    # --------------------------------------------------------
    # AUTO SIZE COLUMNS
    # --------------------------------------------------------

    for column_cells in sheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )


        for cell in column_cells:

            value = cell.value

            if value is None:
                continue


            max_length = max(
                max_length,
                len(str(value))
            )


        sheet.column_dimensions[
            column_letter
        ].width = min(

            max(
                max_length + 2,
                12
            ),

            40
        )


def write_sheet(
    workbook,
    name,
    headers,
    rows
):

    sheet = workbook.create_sheet(
        title=name
    )


    sheet.append(
        headers
    )


    for row in rows:

        sheet.append(
            row
        )


    style_sheet(
        sheet
    )


    return sheet


# ============================================================
# EXPORT DATABASE
# ============================================================

def export_database_excel(
    session: Session
) -> BytesIO:


    workbook = Workbook()


    # Remove the default Excel sheet
    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )


    # ========================================================
    # LOAD PRODUCTS
    # ========================================================

    products = session.exec(
        select(Product)
    ).all()


    product_lookup = {

        product.id: product

        for product in products

        if product.id is not None
    }


    # ========================================================
    # PRODUCTS SHEET
    # ========================================================

    write_sheet(
        workbook,

        "Products",

        [
            "ID",
            "SKU",
            "Brand",
            "Model",
            "Cost Price",
            "Selling Price",
        ],

        [
            [
                product.id,
                product.sku,
                product.brand,
                product.model,
                product.cost_price,
                product.selling_price,
            ]

            for product in products
        ]
    )


    # ========================================================
    # LOAD SUPPLIERS
    # ========================================================

    suppliers = session.exec(
        select(Supplier)
    ).all()


    supplier_lookup = {

        supplier.id: supplier

        for supplier in suppliers

        if supplier.id is not None
    }


    # ========================================================
    # SUPPLIERS SHEET
    # ========================================================

    write_sheet(
        workbook,

        "Suppliers",

        [
            "ID",
            "Name",
            "Phone",
            "Email",
            "Active",
        ],

        [
            [
                supplier.id,
                supplier.name,
                supplier.phone,
                supplier.email,
                supplier.active,
            ]

            for supplier in suppliers
        ]
    )


    # ========================================================
    # LOAD CUSTOMERS
    # ========================================================

    customers = session.exec(
        select(Customer)
    ).all()


    customer_lookup = {

        customer.id: customer

        for customer in customers

        if customer.id is not None
    }


    # ========================================================
    # CUSTOMERS SHEET
    # ========================================================

    write_sheet(
        workbook,

        "Customers",

        [
            "ID",
            "Name",
            "Phone",
            "Email",
            "Active",
        ],

        [
            [
                customer.id,
                customer.name,
                customer.phone,
                customer.email,
                customer.active,
            ]

            for customer in customers
        ]
    )


    # ========================================================
    # LOAD PURCHASES
    # ========================================================

    purchases = session.exec(
        select(Purchase)
    ).all()


    purchase_lookup = {

        purchase.id: purchase

        for purchase in purchases

        if purchase.id is not None
    }


    # ========================================================
    # PURCHASES SHEET
    # ========================================================

    purchase_rows = []


    for purchase in purchases:

        supplier = supplier_lookup.get(
            purchase.supplier_id
        )


        supplier_name = (

            supplier.name

            if supplier

            else "Unknown Supplier"
        )


        purchase_rows.append([

            purchase.id,

            supplier_name,

            purchase.invoice_number,

            format_datetime(
                purchase.purchase_date
            ),

            format_datetime(
                purchase.delivery_date
            ),

            enum_value(
                purchase.status
            ),
        ])


    write_sheet(
        workbook,

        "Purchases",

        [
            "ID",
            "Supplier",
            "Invoice Number",
            "Purchase Date",
            "Delivery Date",
            "Status",
        ],

        purchase_rows
    )


    # ========================================================
    # LOAD PURCHASE ITEMS
    # ========================================================

    purchase_items = session.exec(
        select(PurchaseItem)
    ).all()


    # ========================================================
    # PURCHASE ITEMS SHEET
    # ========================================================

    purchase_item_rows = []


    for item in purchase_items:

        product = product_lookup.get(
            item.product_id
        )


        purchase = purchase_lookup.get(
            item.purchase_id
        )


        # ----------------------------------------------------
        # PRODUCT DISPLAY
        # ----------------------------------------------------

        if product:

            product_name = (
                f"{product.brand} "
                f"{product.model}"
            )

            product_brand = product.brand

            product_model = product.model

            product_sku = product.sku

        else:

            product_name = "Unknown Product"

            product_brand = "Unknown"

            product_model = "Unknown"

            product_sku = "Unknown"


        # ----------------------------------------------------
        # PURCHASE DISPLAY
        # ----------------------------------------------------

        purchase_reference = (

            purchase.invoice_number

            if purchase

            else f"Purchase #{item.purchase_id}"
        )


        line_total = (
            item.quantity
            * item.unit_cost
        )


        purchase_item_rows.append([

            item.id,

            purchase_reference,

            product_name,

            product_brand,

            product_model,

            product_sku,

            item.quantity,

            item.unit_cost,

            line_total,
        ])


    write_sheet(
        workbook,

        "Purchase Items",

        [
            "ID",
            "Purchase",
            "Product",
            "Brand",
            "Model",
            "SKU",
            "Quantity",
            "Unit Cost",
            "Line Total",
        ],

        purchase_item_rows
    )


    # ========================================================
    # LOAD SALES
    # ========================================================

    sales = session.exec(
        select(Sale)
    ).all()


    sale_lookup = {

        sale.id: sale

        for sale in sales

        if sale.id is not None
    }


    # ========================================================
    # SALES SHEET
    # ========================================================

    sale_rows = []


    for sale in sales:

        customer = customer_lookup.get(
            sale.customer_id
        )


        customer_name = (

            customer.name

            if customer

            else "Unknown Customer"
        )


        sale_rows.append([

            sale.id,

            customer_name,

            sale.invoice_number,

            format_datetime(
                sale.sale_date
            ),

            format_datetime(
                sale.delivery_date
            ),

            enum_value(
                sale.sale_status
            ),
        ])


    write_sheet(
        workbook,

        "Sales",

        [
            "ID",
            "Customer",
            "Invoice Number",
            "Sale Date",
            "Delivery Date",
            "Status",
        ],

        sale_rows
    )


    # ========================================================
    # LOAD SALE ITEMS
    # ========================================================

    sale_items = session.exec(
        select(SaleItem)
    ).all()


    # ========================================================
    # SALE ITEMS SHEET
    # ========================================================

    sale_item_rows = []


    for item in sale_items:

        product = product_lookup.get(
            item.product_id
        )


        sale = sale_lookup.get(
            item.sale_id
        )


        # ----------------------------------------------------
        # PRODUCT DISPLAY
        # ----------------------------------------------------

        if product:

            product_name = (
                f"{product.brand} "
                f"{product.model}"
            )

            product_brand = product.brand

            product_model = product.model

            product_sku = product.sku

        else:

            product_name = "Unknown Product"

            product_brand = "Unknown"

            product_model = "Unknown"

            product_sku = "Unknown"


        # ----------------------------------------------------
        # SALE DISPLAY
        # ----------------------------------------------------

        sale_reference = (

            sale.invoice_number

            if sale

            else f"Sale #{item.sale_id}"
        )


        # ----------------------------------------------------
        # CALCULATIONS
        # ----------------------------------------------------

        subtotal = (
            item.quantity
            * item.unit_price
        )


        discount_amount = (
            subtotal
            * item.discount_percentage
            / 100
        )


        total = (
            subtotal
            - discount_amount
        )


        sale_item_rows.append([

            item.id,

            sale_reference,

            product_name,

            product_brand,

            product_model,

            product_sku,

            item.quantity,

            item.unit_price,

            item.discount_percentage,

            subtotal,

            discount_amount,

            total,
        ])


    write_sheet(
        workbook,

        "Sale Items",

        [
            "ID",
            "Sale",
            "Product",
            "Brand",
            "Model",
            "SKU",
            "Quantity",
            "Unit Price",
            "Discount %",
            "Subtotal",
            "Discount Amount",
            "Line Total",
        ],

        sale_item_rows
    )


    # ========================================================
    # LOAD INVENTORY MOVEMENTS
    # ========================================================

    movements = session.exec(
        select(InventoryMovement)
    ).all()


    # ========================================================
    # INVENTORY MOVEMENTS SHEET
    # ========================================================

    movement_rows = []


    for movement in movements:

        product = product_lookup.get(
            movement.product_id
        )


        if product:

            product_name = (
                f"{product.brand} "
                f"{product.model}"
            )

            product_sku = (
                product.sku
            )

        else:

            product_name = (
                "Unknown Product"
            )

            product_sku = (
                "Unknown"
            )


        # ----------------------------------------------------
        # HUMAN-READABLE REFERENCE
        # ----------------------------------------------------

        reference = None


        if (
            movement.reference_type
            and
            movement.reference_id
        ):

            reference_type = (
                movement.reference_type
                .upper()
            )


            if (
                reference_type
                == "SALE"
            ):

                related_sale = sale_lookup.get(
                    movement.reference_id
                )


                if related_sale:

                    reference = (
                        related_sale.invoice_number
                    )


            elif (
                reference_type
                == "PURCHASE"
            ):

                related_purchase = (
                    purchase_lookup.get(
                        movement.reference_id
                    )
                )


                if related_purchase:

                    reference = (
                        related_purchase.invoice_number
                    )


        if reference is None:

            reference = (
                str(
                    movement.reference_id
                )

                if movement.reference_id
                is not None

                else ""
            )


        movement_rows.append([

            movement.id,

            product_name,

            product_sku,

            enum_value(
                movement.movement_type
            ),

            movement.quantity,

            format_datetime(
                movement.created_at
            ),

            movement.reference_type,

            reference,
        ])


    write_sheet(
        workbook,

        "Inventory Movements",

        [
            "ID",
            "Product",
            "SKU",
            "Movement Type",
            "Quantity",
            "Created At",
            "Reference Type",
            "Reference",
        ],

        movement_rows
    )


    # ========================================================
    # SAVE INTO MEMORY
    # ========================================================

    output = BytesIO()


    workbook.save(
        output
    )


    output.seek(0)


    return output